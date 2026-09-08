"""Render a `reports.Report` as CSV, XLSX, or PDF bytes (M8, PRD §18.5)."""

from __future__ import annotations

import csv
import io

from app.services.reports import Report

CONTENT_TYPES = {
    "csv": "text/csv",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


def _cells(report: Report) -> tuple[list[str], list[list[str]]]:
    headers = [c["label"] for c in report.columns]
    keys = [c["key"] for c in report.columns]
    body = [[_str(row.get(k, "")) for k in keys] for row in report.rows]
    return headers, body


def _str(v) -> str:
    if v is None:
        return ""
    return str(v)


def to_csv(report: Report) -> bytes:
    headers, body = _cells(report)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(headers)
    w.writerows(body)
    if report.summary:
        w.writerow([])
        for k, v in report.summary.items():
            w.writerow([k.replace("_", " ").title(), _str(v)])
    return buf.getvalue().encode("utf-8-sig")  # BOM so Excel opens UTF-8 cleanly


def to_xlsx(report: Report) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    headers, body = _cells(report)
    wb = Workbook()
    ws = wb.active
    ws.title = report.label[:31]

    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in body:
        ws.append(row)

    if report.summary:
        ws.append([])
        for k, v in report.summary.items():
            ws.append([k.replace("_", " ").title(), _str(v)])

    for col in ws.columns:
        width = max((len(_str(c.value)) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(width + 2, 48)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(report: Report) -> bytes:
    from fpdf import FPDF

    headers, body = _cells(report)

    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 8, report.label, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 8)
    meta = f"Generated for {report.generated_for}"
    if report.filters:
        meta += "  ·  " + ", ".join(f"{k}={v}" for k, v in report.filters.items())
    pdf.set_text_color(110, 110, 110)
    pdf.cell(0, 5, meta, new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(0, 0, 0)
    pdf.ln(2)

    usable = pdf.w - pdf.l_margin - pdf.r_margin
    col_w = usable / max(len(headers), 1)

    pdf.set_font("Helvetica", "B", 8)
    pdf.set_fill_color(238, 238, 238)
    for h in headers:
        pdf.cell(col_w, 7, _clip(h, col_w, pdf), border=1, fill=True)
    pdf.ln()

    pdf.set_font("Helvetica", "", 8)
    for row in body:
        for value in row:
            pdf.cell(col_w, 6, _clip(value, col_w, pdf), border=1)
        pdf.ln()

    if report.summary:
        pdf.ln(3)
        pdf.set_font("Helvetica", "B", 9)
        summary = "   ".join(
            f"{k.replace('_', ' ').title()}: {_str(v)}" for k, v in report.summary.items()
        )
        pdf.cell(0, 6, summary, new_x="LMARGIN", new_y="NEXT")

    out = pdf.output()
    return bytes(out)


def _clip(text: str, width_mm: float, pdf) -> str:
    text = _str(text)
    while text and pdf.get_string_width(text) > width_mm - 2:
        text = text[:-1]
    return text


def render(report: Report, fmt: str) -> bytes:
    return {"csv": to_csv, "xlsx": to_xlsx, "pdf": to_pdf}[fmt](report)
